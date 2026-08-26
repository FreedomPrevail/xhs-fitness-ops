import json
import tempfile
import unittest
import urllib.parse
from pathlib import Path

from PIL import Image
from pptx import Presentation

from canva_connect.client import CanvaClient, CanvaConfig, make_pkce_pair
from content.carousel.pptx_builder import ensure_editable_pptx


class CanvaConnectTests(unittest.TestCase):
    def setUp(self):
        self.config = CanvaConfig("client-id", "client-secret", "http://127.0.0.1:8765/oauth/callback")

    def test_pkce_and_authorization_url(self):
        verifier, challenge = make_pkce_pair()
        self.assertGreaterEqual(len(verifier), 43)
        self.assertTrue(challenge)
        url = CanvaClient(self.config).authorization_url(challenge, "state-1")
        query = urllib.parse.parse_qs(urllib.parse.urlparse(url).query)
        self.assertEqual(query["code_challenge_method"], ["s256"])
        self.assertEqual(query["scope"], ["design:content:write"])
        self.assertEqual(query["redirect_uri"], [self.config.redirect_uri])

    def test_import_design_polls_and_returns_edit_url(self):
        calls = []

        def transport(method, url, headers, data, timeout):
            calls.append((method, url, headers, data))
            if method == "POST":
                return {"job": {"id": "job-1", "status": "in_progress"}}
            return {"job": {"id": "job-1", "status": "success", "result": {"designs": [{
                "id": "design-1", "title": "测试", "page_count": 2,
                "urls": {"edit_url": "https://www.canva.com/design/edit", "view_url": "https://www.canva.com/design/view"},
            }]}}}

        with tempfile.TemporaryDirectory() as tmp:
            deck = Path(tmp) / "sample.pptx"
            deck.write_bytes(b"pptx-bytes")
            client = CanvaClient(self.config, transport=transport)
            client.access_token = lambda: "access-token"  # type: ignore[method-assign]
            result = client.import_design(deck, title="测试", poll_interval=0.2)
        self.assertTrue(result["ok"])
        self.assertEqual(result["designs"][0]["id"], "design-1")
        self.assertEqual(calls[0][0], "POST")
        metadata = json.loads(calls[0][2]["Import-Metadata"])
        self.assertEqual(metadata["mime_type"], "application/vnd.openxmlformats-officedocument.presentationml.presentation")

    def test_layered_pptx_contains_editable_objects(self):
        with tempfile.TemporaryDirectory() as tmp:
            package = Path(tmp) / "xhs_20260825_000001"
            assets = package / "source_assets"
            assets.mkdir(parents=True)
            image = assets / "xhs_20260825_000001_cover_01.png"
            Image.new("RGB", (400, 500), "#F6C7AE").save(image)
            manifest = {
                "post_id": package.name,
                "selected_cover_index": 0,
                "source_assets": [{"path": f"source_assets/{image.name}", "page_no": 1}],
                "pages": [{"page_no": 1, "title": "封面"}, {"page_no": 2, "title": "内容"}],
            }
            layers = {"pages": [
                {"page_no": 1, "title": "女生新手健身", "subtitle": "先学会这一套", "badge": "新手指南", "modules": []},
                {"page_no": 2, "title": "训练前检查", "subtitle": "不会调节就问", "badge": "安全", "modules": [{
                    "title": "三项检查", "type": "action_cards", "items": [
                        {"label": "插销", "detail": "确认完全插入"},
                        {"label": "锁扣", "detail": "确认已经闭合"},
                    ],
                }]},
            ]}
            (package / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
            (package / "editable_text_layers.json").write_text(json.dumps(layers, ensure_ascii=False), encoding="utf-8")
            deck = ensure_editable_pptx(package)
            prs = Presentation(deck)
            self.assertEqual(len(prs.slides), 2)
            self.assertGreater(sum(1 for shape in prs.slides[0].shapes if shape.has_text_frame), 3)
            self.assertGreater(sum(1 for shape in prs.slides[0].shapes if shape.shape_type == 13), 0)
            saved_manifest = json.loads((package / "manifest.json").read_text(encoding="utf-8"))
            self.assertEqual(saved_manifest["canva_import_mode"], "connect_api_pptx")


if __name__ == "__main__":
    unittest.main()
